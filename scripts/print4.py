import os
import sys
from docx import Document
import win32com.client as win32

def edit_word_doc(value_dict):
    # Open the document template
    doc = Document('./output2.docx')
    
    # Replace placeholders with actual values
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for key, value in value_dict.items():
                    if key in cell.text:
                        cell.text = cell.text.replace(key, value)

    
    # Save the edited document
    output_path = os.path.abspath('output2.docx')
    doc.save(output_path)
    
    # Print the edited document
    print_word_doc(output_path)

def print_word_doc(doc_path):
    # Create a Word application instance
    word_app = win32.Dispatch('Word.Application')
    
    # Open the document
    doc = word_app.Documents.Open(doc_path)
    
    # Print the document
    doc.PrintOut()
    
    # Close the document without saving
    doc.Close(False)
    
    # Quit the Word application
    word_app.Quit()

if __name__ == "__main__":
    # Create a dictionary from command line arguments
    value_dict = {
        '{REGNO}': sys.argv[1],
        '{QUAL}': sys.argv[2],
        '{OCCUP}': sys.argv[3],
        '{MARITIAL}': sys.argv[4],
        '{DOB}': sys.argv[5],
        '{EMAIL}': sys.argv[6],
        '{CASEBY}':sys.argv[7],
        '{REFERRED}': sys.argv[8],
        '{DATE}': sys.argv[9],
    }
    
    # Edit and print the Word document
    edit_word_doc(value_dict)





#   TextEditingController _additionalDescriptionController =
#       TextEditingController();
#   TextEditingController _additionalNotesController = TextEditingController();
#   TextEditingController _followUpDateController = TextEditingController();
#   TextEditingController _courseSuggestedController = TextEditingController();

#   TextEditingController _caseTakenByController = TextEditingController();

